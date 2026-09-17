"""MVP2 batch acceptance against the immutable report API and database."""

import argparse
import hashlib
import json
import os
import urllib.parse
import urllib.request
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pymysql


parser = argparse.ArgumentParser()
parser.add_argument("--batch", required=True)
parser.add_argument("--database", required=True)
parser.add_argument("--port", type=int, default=18096)
parser.add_argument("--output", required=True)
args = parser.parse_args()

batch = args.batch
root = Path(__file__).with_name(args.output)
root.mkdir(exist_ok=False)
headers = {
    "x-company-user-id": "internal-acceptance",
    "x-company-user-name": "Internal%20Acceptance",
    "x-company-role": "viewer",
}


def fetch(route, params, require_batch=True):
    query = urllib.parse.urlencode({"batch": batch, **params})
    url = f"http://127.0.0.1:{args.port}/aftersales/{route}?{query}"
    request = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(request, timeout=180) as response:
        body = response.read()
        response_headers = {key.lower(): value for key, value in response.headers.items()}
    if require_batch:
        assert response_headers.get("x-data-batch") == batch, (url, response_headers)
    return body, response_headers


def pages(route, params):
    rows = []
    total = None
    first_payload = None
    while total is None or len(rows) < total:
        body, _ = fetch(route, {**params, "offset": len(rows), "limit": 200})
        payload = json.loads(body)
        if first_payload is None:
            first_payload = payload
        total = payload.get("rowGroupCount", payload.get("total"))
        assert isinstance(total, int)
        assert payload["offset"] == len(rows)
        if not payload["rows"] and len(rows) < total:
            raise AssertionError(f"early empty page for {route}: {len(rows)} < {total}")
        rows.extend(payload["rows"])
    assert len(rows) == total
    return rows, first_payload


def export(name, params):
    body, response_headers = fetch("export", params)
    path = root / f"{name}.xlsx"
    path.write_bytes(body)
    path.chmod(0o600)
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    worksheet_rows = worksheet.max_row
    labels = [worksheet.cell(row=index, column=1).value for index in range(1, min(20, worksheet.max_row) + 1)]
    workbook.close()
    return {
        "path": str(path),
        "sha256": hashlib.sha256(body).hexdigest(),
        "rowCountHeader": response_headers.get("x-export-row-count"),
        "buildId": response_headers.get("x-build-id"),
        "worksheetRows": worksheet_rows,
        "labels": labels,
    }


connection = pymysql.connect(
    host="127.0.0.1",
    user="root",
    database=args.database,
    cursorclass=pymysql.cursors.DictCursor,
)
checks = []
try:
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT summary_json FROM crm_report_batch WHERE batch_id=%s AND status IN ('candidate','verified')",
            (batch,),
        )
        batch_row = cursor.fetchone()
        assert batch_row
        summary = json.loads(batch_row["summary_json"])

    assert summary["issueCoverageApproved"] is True
    assert summary["salesCoverageApproved"] is False
    conservation = summary["sourceInputConservation"]
    assert conservation["sourceInputs"] == conservation["canonicalIssues"] + conservation["mergedDuplicates"]
    assert conservation["sourceInputs"] == sum(conservation["bySource"].values())

    ready_body, ready_headers = fetch("readyz", {}, require_batch=False)
    ready = json.loads(ready_body)
    assert ready["ok"] is True
    assert ready["analytics"] == {
        "state": "issue_facts_ready",
        "code": "denominators_pending_verification",
    }

    for selected in summary["releaseScope"]["selectableReleaseWeeks"]:
        start = date.fromisoformat(selected)
        end = start + timedelta(days=7)
        products, products_payload = pages("api/products", {"week_start": selected})
        details, detail_payload = pages(
            "api/issues", {"basis": "created", "scope": "all", "week_start": selected}
        )
        audit = detail_payload["auditSummary"]

        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT COUNT(*) total_issues,
                       SUM(included_operating=1) operating_issues,
                       SUM(included_operating=0 AND TRIM(problem1) IN ('快递问题','库房问题','买家问题','产品问题','运营问题')) excluded_issues,
                       SUM(TRIM(problem1) NOT IN ('快递问题','库房问题','买家问题','产品问题','运营问题')) unclassified_issues,
                       SUM(included_operating=1 AND merchant_code<>'') matched_operating_issues,
                       SUM(included_operating=1 AND merchant_code='') unmatched_operating_issues
                FROM crm_report_issue
                WHERE batch_id=%s AND created_at>=%s AND created_at<%s
                """,
                (batch, start, end),
            )
            expected = cursor.fetchone()
            cursor.execute(
                """
                SELECT source_system, COUNT(*) issues
                FROM crm_report_issue
                WHERE batch_id=%s AND created_at>=%s AND created_at<%s
                GROUP BY source_system ORDER BY source_system
                """,
                (batch, start, end),
            )
            expected_sources = {row["source_system"]: row["issues"] for row in cursor.fetchall()}

        expected_audit = {
            "totalIssues": expected["total_issues"],
            "operatingIssues": int(expected["operating_issues"] or 0),
            "excludedIssues": int(expected["excluded_issues"] or 0),
            "unclassifiedIssues": int(expected["unclassified_issues"] or 0),
            "matchedOperatingIssues": int(expected["matched_operating_issues"] or 0),
            "unmatchedOperatingIssues": int(expected["unmatched_operating_issues"] or 0),
        }
        actual_audit = {key: audit[key] for key in expected_audit}
        assert actual_audit == expected_audit, {"week": selected, "actual": actual_audit, "expected": expected_audit}
        assert {row["sourceSystem"]: row["issues"] for row in audit["sourceInputs"]} == expected_sources
        assert audit["operatingIssues"] + audit["excludedIssues"] + audit["unclassifiedIssues"] == audit["totalIssues"]
        assert audit["matchedOperatingIssues"] + audit["unmatchedOperatingIssues"] == audit["operatingIssues"]

        assert sum(row["currentIssues"] for row in details) == audit["totalIssues"]
        status_totals = Counter()
        matched_by_spec = Counter()
        for row in details:
            status_totals[row["classificationStatus"]] += row["currentIssues"]
            if row["includedInOperating"] and row["code"] != "未填":
                matched_by_spec[row["code"]] += row["currentIssues"]
            for week in row["weeks"]:
                assert week["orders"] is None and week["sales"] is None and week["rate"] is None
        assert status_totals["included"] == audit["operatingIssues"]
        assert status_totals["excluded"] == audit["excludedIssues"]
        assert status_totals["unclassified"] == audit["unclassifiedIssues"]
        assert sum(row["selectedIssues"] for row in products) == audit["matchedOperatingIssues"]
        assert {row["code"]: row["selectedIssues"] for row in products} == dict(matched_by_spec)
        assert all("scope=all" in row["detailHref"] for row in products)

        exports = {
            "products": export(
                f"{selected}-products",
                {"view": "products", "basis": "created", "week_start": selected},
            ),
            "details": export(
                f"{selected}-details",
                {"view": "detail", "basis": "created", "scope": "all", "week_start": selected},
            ),
        }
        assert "待匹配记录" not in exports["details"]["labels"]
        assert exports["products"]["buildId"] == ready_headers.get("x-build-id")
        assert exports["details"]["buildId"] == ready_headers.get("x-build-id")

        check = {
            "weekStart": selected,
            "productRows": len(products),
            "detailRowGroups": len(details),
            "audit": {key: audit[key] for key in expected_audit},
            "sources": audit["sourceInputs"],
            "exports": exports,
        }
        checks.append(check)
        print(json.dumps(check, ensure_ascii=False), flush=True)

    result = {
        "batch": batch,
        "buildId": ready_headers.get("x-build-id"),
        "readyCurrentBatchBeforePublication": ready["source"]["batchId"],
        "issueCoverageApproved": True,
        "salesCoverageApproved": False,
        "sourceInputConservation": conservation,
        "checks": checks,
        "productionPublished": False,
        "employeeEntryVerified": False,
    }
    result_path = root / "summary.json"
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    result_path.chmod(0o600)
finally:
    connection.rollback()
    connection.close()

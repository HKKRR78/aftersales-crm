#!/usr/bin/env python3
from __future__ import annotations

import base64
import csv
import fcntl
import hashlib
import hmac
import html
import io
import ipaddress
import json
import os
import secrets
import sys
import time
from contextlib import contextmanager
from datetime import datetime
from http import HTTPStatus
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlencode, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEPLOY_ROOT = Path(os.environ.get("ECOM_DEPLOY_ROOT", "/Users/Shared/ecom-profit"))
APP_ROOT = Path(os.environ.get("ECOM_PROJECT_ROOT", DEPLOY_ROOT / "app"))
RUNTIME_ROOT = Path(os.environ.get("CRM_RUNTIME_ROOT", DEPLOY_ROOT / "projects" / "crm" / "runtime" / "crm_dashboard"))
REPORT_ROOT = Path(os.environ.get("AFTERSALES_REPORT_OUT_ROOT", DEPLOY_ROOT / "outputs" / "售后经营异常周报"))
PROJECT_REPORT_ROOT = DEPLOY_ROOT / "projects" / "crm" / "outputs" / "售后经营异常周报"
HOST = os.environ.get("CRM_HOST", "0.0.0.0")
PORT = int(os.environ.get("CRM_PORT", "8088"))
SESSION_TTL = int(os.environ.get("CRM_SESSION_TTL_SECONDS", str(7 * 24 * 3600)))
XLSX_EXPORT_ENABLED = os.environ.get("CRM_XLSX_EXPORT_ENABLED", "1").lower() not in {"0", "false", "no"}

USER_FILE = RUNTIME_ROOT / "users.json"
SESSION_FILE = RUNTIME_ROOT / "sessions.json"
INITIAL_PASSWORD_FILE = RUNTIME_ROOT / "initial_admin_password.txt"

ROLES = {
    "viewer": "只读访问",
    "exporter": "导出权限",
    "admin": "管理员",
}

DATA_CACHE = {"path": "", "mtime": 0.0, "data": None}
MYSQL_DATA_CACHE = {"loaded_at": {}, "data": {}, "error": {}}
INTRANET_VIEWER = {
    "username": "intranet_viewer",
    "display_name": "内网访客",
    "role": "viewer",
    "is_active": True,
}
INTRANET_NETWORKS = [
    ipaddress.ip_network("127.0.0.0/8"),
    ipaddress.ip_network("::1/128"),
    ipaddress.ip_network("192.168.0.0/16"),
    ipaddress.ip_network("10.0.0.0/8"),
    ipaddress.ip_network("172.16.0.0/12"),
]
TRUSTED_PROXY_NETWORKS = [
    ipaddress.ip_network(item.strip())
    for item in os.environ.get("CRM_TRUSTED_PROXY_NETWORKS", "127.0.0.0/8,::1/128").split(",")
    if item.strip()
]


def ensure_runtime() -> None:
    RUNTIME_ROOT.mkdir(parents=True, exist_ok=True)
    os.chmod(RUNTIME_ROOT, 0o700)
    if not USER_FILE.exists():
        password = secrets.token_urlsafe(12)
        users = {
            "admin": {
                "username": "admin",
                "display_name": "管理员",
                "role": "admin",
                "is_active": True,
                "password": hash_password(password),
                "created_at": now_text(),
                "updated_at": now_text(),
                "last_login_at": "",
            }
        }
        write_json(USER_FILE, users)
        INITIAL_PASSWORD_FILE.write_text(f"username: admin\npassword: {password}\n", encoding="utf-8")
        os.chmod(INITIAL_PASSWORD_FILE, 0o600)
    if not SESSION_FILE.exists():
        write_json(SESSION_FILE, {})
    for path in (USER_FILE, SESSION_FILE, INITIAL_PASSWORD_FILE):
        if path.exists():
            os.chmod(path, 0o600)


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def read_json(path: Path, fallback):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return fallback


@contextmanager
def json_file_lock(path: Path):
    lock_path = path.with_name(path.name + ".lock")
    descriptor = os.open(lock_path, os.O_CREAT | os.O_RDWR, 0o600)
    os.chmod(lock_path, 0o600)
    with os.fdopen(descriptor, "r+") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def write_json_unlocked(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.chmod(tmp, 0o600)
    tmp.replace(path)
    os.chmod(path, 0o600)


def write_json(path: Path, payload) -> None:
    with json_file_lock(path):
        write_json_unlocked(path, payload)


def mutate_json(path: Path, fallback, mutator):
    with json_file_lock(path):
        payload = read_json(path, fallback)
        result = mutator(payload)
        write_json_unlocked(path, payload)
        return result


def hash_password(password: str, *, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 200_000)
    return f"pbkdf2_sha256$200000${salt}${base64.b64encode(digest).decode('ascii')}"


def verify_password(password: str, encoded: str) -> bool:
    try:
        algo, rounds, salt, digest = encoded.split("$", 3)
        if algo != "pbkdf2_sha256":
            return False
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), int(rounds))
        return hmac.compare_digest(base64.b64encode(actual).decode("ascii"), digest)
    except Exception:
        return False


def latest_dashboard_path() -> Path | None:
    roots = [REPORT_ROOT, PROJECT_REPORT_ROOT]
    candidates = []
    for root in roots:
        if root.exists():
            candidates.extend(root.rglob("售后经营异常周报看板_*.html"))
    if not candidates:
        return None
    return sorted(candidates, key=lambda p: p.stat().st_mtime, reverse=True)[0]


def extract_dashboard_data(path: Path) -> dict:
    mtime = path.stat().st_mtime
    if DATA_CACHE["path"] == str(path) and DATA_CACHE["mtime"] == mtime and DATA_CACHE["data"] is not None:
        return DATA_CACHE["data"]
    text = path.read_text(encoding="utf-8")
    marker = "const DATA = "
    start = text.find(marker)
    if start < 0:
        raise RuntimeError("dashboard DATA not found")
    start += len(marker)
    end_marker = "\nconst DISPLAY_WEEKS"
    end = text.find(end_marker, start)
    if end < 0:
        end = text.find(";\n", start)
    raw = text[start:end].strip().rstrip(";")
    data = json.loads(raw)
    DATA_CACHE.update({"path": str(path), "mtime": mtime, "data": data})
    return data


def normalize_period(value) -> str:
    if isinstance(value, dict):
        value = (value.get("period") or ["closed"])[0]
    return "progress" if str(value or "").lower() == "progress" else "closed"


def load_data(period: str = "closed") -> tuple[dict, Path]:
    period = normalize_period(period)
    mysql_data = load_mysql_data(period)
    if mysql_data:
        return mysql_data, Path("MySQL: ecom_profit.ods_banniu_aftersales")
    if period == "progress":
        return {
            "weeks": [], "orderWeeks": [], "rows": [], "warehouseRows": [], "kpis": {},
            "periodMode": "progress", "periodMeta": {},
            "source": {"sourceType": "mysql_unavailable", "currentPeriod": "本周同进度数据暂不可用"},
        }, Path("MySQL unavailable")
    path = latest_dashboard_path()
    if not path:
        return {"weeks": [], "rows": [], "warehouseRows": [], "kpis": {}, "source": ""}, Path("")
    return extract_dashboard_data(path), path


def load_mysql_data(period: str = "closed") -> dict | None:
    period = normalize_period(period)
    ttl = int(os.environ.get("CRM_MYSQL_CACHE_SECONDS", "300"))
    now = time.time()
    cached = MYSQL_DATA_CACHE["data"].get(period)
    loaded_at = float(MYSQL_DATA_CACHE["loaded_at"].get(period) or 0)
    if cached is not None and now - loaded_at < ttl:
        return cached
    try:
        from crm_mysql_data import load_mysql_dashboard_data

        data = load_mysql_dashboard_data(period=period)
        if data.get("rows") or data.get("warehouseRows") or data.get("periodMode") == "progress":
            MYSQL_DATA_CACHE["loaded_at"][period] = now
            MYSQL_DATA_CACHE["data"][period] = data
            MYSQL_DATA_CACHE["error"][period] = ""
            return data
    except Exception as exc:
        MYSQL_DATA_CACHE["loaded_at"][period] = now
        MYSQL_DATA_CACHE["data"].pop(period, None)
        MYSQL_DATA_CACHE["error"][period] = str(exc)
        sys.stderr.write(f"{now_text()} mysql dashboard fallback: {exc}\n")
    return None


def esc(value) -> str:
    return html.escape(str(value if value is not None else ""))


def fmt_num(value) -> str:
    try:
        return f"{int(float(value or 0)):,}"
    except Exception:
        return "0"


def pct(value) -> str:
    if value in (None, ""):
        return ""
    try:
        return f"{float(value):.2f}%"
    except Exception:
        return esc(value)


def rate_num(value) -> float:
    try:
        return float(value)
    except Exception:
        return 0.0


def week_rate_text(week: dict) -> str:
    if week.get("rateText") not in (None, ""):
        return str(week.get("rateText"))
    if week.get("rate") in (None, ""):
        return ""
    return pct(week.get("rate"))


def week_wow_text(week: dict) -> str:
    if week.get("wowText") not in (None, ""):
        return str(week.get("wowText"))
    if week.get("wow") in (None, ""):
        return ""
    return pct(week.get("wow"))


def heat_class(week: dict) -> str:
    high_rate = rate_num(week.get("rate")) >= 2.0
    rising = rate_num(week.get("wow")) > 0
    if high_rate and rising:
        return "heat-dark"
    if high_rate:
        return "heat-red"
    if rising:
        return "heat-yellow"
    return ""


def has_recent_rise(row: dict) -> bool:
    return any(rate_num(w.get("wow")) > 0 for w in row.get("weeks") or [])


def apply_comparison_fields(row: dict) -> dict:
    weeks = row.get("weeks") or []
    previous = weeks[-2] if len(weeks) >= 2 else {}
    current = weeks[-1] if weeks else {}
    previous_rate = previous.get("rate")
    current_rate = current.get("rate")
    delta_pp = current_rate - previous_rate if current_rate is not None and previous_rate is not None else None
    expected = previous_rate / 100 * int(current.get("orders") or 0) if previous_rate is not None else None
    excess = int(current.get("issues") or 0) - expected if expected is not None else None
    min_excess = float(os.environ.get("CRM_PROGRESS_ALERT_MIN_EXCESS", "3"))
    min_delta_pp = float(os.environ.get("CRM_PROGRESS_ALERT_MIN_DELTA_PP", "0.3"))
    if delta_pp is None:
        level, rank = "灰色", 0
    elif excess is not None and excess >= min_excess and delta_pp >= min_delta_pp:
        level, rank = "红色", 3
    elif delta_pp > 0:
        level, rank = "黄色", 2
    elif delta_pp < 0:
        level, rank = "绿色", 1
    else:
        level, rank = "灰色", 0
    row.update({
        "latestIssues": int(current.get("issues") or 0),
        "latestRate": current_rate,
        "latestWow": current.get("wow"),
        "latestWowText": current.get("wowText") or "",
        "previousIssues": int(previous.get("issues") or 0),
        "previousOrders": int(previous.get("orders") or 0),
        "previousRate": previous_rate,
        "currentIssues": int(current.get("issues") or 0),
        "currentOrders": int(current.get("orders") or 0),
        "currentRate": current_rate,
        "rateDeltaPp": delta_pp,
        "relativeChange": current.get("wow"),
        "expectedIssues": expected,
        "excessIssues": excess,
        "alertLevel": level,
        "alertRank": rank,
    })
    return row


def source_info(data: dict) -> dict:
    source = data.get("source") or {}
    if isinstance(source, dict):
        return source
    return {"currentPeriod": str(source), "orderKey": "", "aftersalesKey": ""}


def role_label(role: str) -> str:
    return ROLES.get(role, role)


def can_export(user: dict) -> bool:
    return user.get("role") in {"exporter", "admin"}


def can_export_excel(user: dict) -> bool:
    return XLSX_EXPORT_ENABLED and user.get("username") != INTRANET_VIEWER["username"] and user.get("role") in ROLES


def is_admin(user: dict) -> bool:
    return user.get("role") == "admin"


def excel_export_link(view: str, user: dict, query: dict | None = None) -> str:
    if not can_export_excel(user):
        return ""
    params = [("view", view), ("format", "xlsx")]
    for key in ("q", "merchant_code", "doudian_product_id", "p1", "p2", "rise", "period"):
        value = ((query or {}).get(key) or [""])[0]
        if value:
            params.append((key, value))
    return f'<a class="btn" href="/crm/export?{esc(urlencode(params))}">导出 Excel</a>'


def product_detail_url(code: str, period: str = "closed") -> str:
    code = str(code or "").strip()
    if not code or code == "未填":
        return ""
    return "/crm/dashboard/detail?" + urlencode({"merchant_code": code, "period": normalize_period(period)})


def layout(title: str, body: str, user: dict | None = None, period: str = "closed") -> bytes:
    nav = ""
    if user:
        period = normalize_period(period)
        suffix = "?period=progress" if period == "progress" else ""
        export_query = "&period=progress" if period == "progress" else ""
        export_link = f'<a href="/crm/export?view=detail{export_query}">导出明细</a>' if can_export(user) else ""
        admin_link = '<a href="/crm/admin/users">用户管理</a>' if is_admin(user) else ""
        nav = f"""
        <nav>
          <a href="/crm">首页</a>
          <a href="/crm/dashboard{suffix}">日更售后 Dashboard</a>
          <a href="/crm/dashboard/detail{suffix}">售后明细</a>
          <a href="/crm/warehouse{suffix}">仓库问题视图</a>
          <a href="/crm/products{suffix}">商品问题视图</a>
          {export_link}
          {admin_link}
          <span class="who">{esc(user.get('display_name') or user.get('username'))} · {esc(role_label(user.get('role','')))}</span>
          <a href="/crm/logout">退出</a>
        </nav>
        """
    page = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)}</title>
<style>
body{{margin:0;font-family:"Microsoft YaHei","PingFang SC",Arial,sans-serif;background:#f5f7f6;color:#1d2723}}
header{{background:#164735;color:#fff;padding:18px 24px}}
h1{{margin:0;font-size:22px}} .sub{{color:#dcebe4;font-size:12px;margin-top:6px}}
nav{{display:flex;gap:12px;align-items:center;flex-wrap:wrap;padding:10px 24px;background:#fff;border-bottom:1px solid #dce5df;position:sticky;top:0;z-index:10}}
nav a{{color:#164735;text-decoration:none;font-weight:700;font-size:13px}} .who{{margin-left:auto;color:#66766f;font-size:12px}}
main{{padding:20px 24px;max-width:1800px;margin:0 auto}} .grid{{display:grid;grid-template-columns:repeat(4,minmax(160px,1fr));gap:12px;margin-bottom:16px}}
.card{{background:#fff;border:1px solid #dce5df;border-radius:10px;padding:14px;box-shadow:0 6px 18px rgba(25,60,45,.05)}}
.label{{font-size:12px;color:#68766f}} .value{{font-size:26px;font-weight:850;color:#164735;margin-top:6px}}
.toolbar{{display:flex;gap:10px;flex-wrap:wrap;margin:0 0 12px}} input,select,button{{font-family:inherit;border:1px solid #dce5df;border-radius:8px;padding:9px 10px;background:#fff}}
button,.btn{{background:#164735;color:#fff;border:0;text-decoration:none;display:inline-block}} .muted{{color:#68766f;font-size:12px}}
table{{border-collapse:separate;border-spacing:0;width:100%;min-width:max-content;font-size:12px;background:#fff;border:1px solid #dce5df;border-radius:10px;overflow:visible}}
thead th{{background:#164735;color:#fff;text-align:left;padding:8px;white-space:nowrap;position:sticky;top:44px;z-index:20}} td{{border-top:1px solid #edf1ee;padding:8px;vertical-align:top}}
.num{{text-align:right;font-variant-numeric:tabular-nums}} .bad{{color:#b33f35;font-weight:800}} .ok{{color:#164735;font-weight:800}}
.table-wrap{{overflow:visible;border-radius:10px}} .section-title{{display:flex;align-items:end;justify-content:space-between;gap:12px;margin:18px 0 10px}}
.section-actions{{display:flex;align-items:center;gap:8px;flex-wrap:wrap}}
.section-title h2{{margin:0;font-size:18px}} .stack{{display:flex;flex-direction:column;gap:16px}}
.pill{{display:inline-block;border-radius:999px;padding:2px 8px;background:#eaf2ee;color:#164735;font-weight:700;font-size:12px}}
.heat-red{{background:#f7d9d4;color:#8f2118;font-weight:800}} .heat-yellow{{background:#fff1b8;color:#6e5200;font-weight:800}}
.heat-dark{{background:#8f2118;color:#fff;font-weight:900}} .week-cell{{min-width:104px;line-height:1.55}} .base-col{{min-width:92px}}
.period-tabs{{display:flex;gap:8px;margin:0 0 16px}} .period-tabs .btn{{background:#dfe9e4;color:#164735}}
.period-tabs .active{{background:#164735;color:#fff}} .alert-red{{color:#8f2118;font-weight:900}}
.alert-yellow{{color:#8a6500;font-weight:850}} .alert-green{{color:#16734b;font-weight:850}} .alert-gray{{color:#68766f;font-weight:700}}
.wide-name{{min-width:180px}} .source{{word-break:break-all}}
.login{{max-width:420px;margin:60px auto}} iframe{{width:100%;height:calc(100vh - 128px);border:1px solid #dce5df;border-radius:10px;background:#fff}}
@media(max-width:900px){{.grid{{grid-template-columns:1fr 1fr}} nav .who{{margin-left:0}}}}
</style>
</head>
<body><header><h1>{esc(title)}</h1><div class="sub">内部 CRM · 售后经营异常</div></header>{nav}<main>{body}</main></body></html>"""
    return page.encode("utf-8")


def row_matches(row: dict, q: str) -> bool:
    if not q:
        return True
    return q.lower() in json.dumps(row, ensure_ascii=False).lower()


def filtered_rows(rows: list[dict], query: dict) -> list[dict]:
    q = (query.get("q") or [""])[0].strip()
    merchant_code = (query.get("merchant_code") or [""])[0].strip()
    doudian_product_id = (query.get("doudian_product_id") or [""])[0].strip()
    p1 = (query.get("p1") or [""])[0].strip()
    p2 = (query.get("p2") or [""])[0].strip()
    rise_only = (query.get("rise") or [""])[0].strip() == "1"
    out = []
    for row in rows:
        if merchant_code and merchant_code != str(row.get("code") or "").strip():
            continue
        if p1 and row.get("p1") != p1:
            continue
        if p2 and row.get("p2") != p2:
            continue
        if doudian_product_id and doudian_product_id != str(row.get("doudian_product_id") or "").strip():
            continue
        if not row_matches(row, q):
            continue
        if rise_only and not has_recent_rise(row):
            continue
        out.append(row)
    return out


def aggregate_products(rows: list[dict], weeks: list[str], period: str = "closed") -> list[dict]:
    groups: dict[str, dict] = {}
    for row in rows:
        code = str(row.get("code") or "")
        name = str(row.get("name") or "")
        g = groups.setdefault(code, {"code": code, "name": name, "doudian_product_id": "", "_doudian_product_ids": set(), "latestIssues": 0, "total5": 0, "weeks": [{"label": w, "issues": 0, "orders": 0, "rate": 0.0, "wow": None, "wowText": ""} for w in weeks]})
        doudian_product_id = str(row.get("doudian_product_id") or "").strip()
        if doudian_product_id:
            g["_doudian_product_ids"].add(doudian_product_id)
        if name and (not g["name"] or len(name) > len(str(g["name"]))):
            g["name"] = name
        g["latestIssues"] += int(row.get("latestIssues") or 0)
        g["total5"] += int(row.get("total5") or 0)
        for i, w in enumerate(row.get("weeks") or []):
            g["weeks"][i]["issues"] += int(w.get("issues") or 0)
            g["weeks"][i]["orders"] = max(g["weeks"][i]["orders"], int(w.get("orders") or 0))
    for g in groups.values():
        prev_rate = None
        for w in g["weeks"]:
            w["rate"] = w["issues"] / w["orders"] * 100 if w["orders"] else None
            w["rateText"] = pct(w["rate"]) if w["orders"] else ""
            if w["rate"] is None or prev_rate is None:
                w["wow"] = None
                w["wowText"] = ""
            elif prev_rate == 0 and w["rate"] > 0:
                w["wow"] = 9999.0
                w["wowText"] = "上期为0"
            elif prev_rate == 0:
                w["wow"] = 0.0
                w["wowText"] = "0.00%"
            else:
                w["wow"] = (w["rate"] - prev_rate) / prev_rate * 100
                w["wowText"] = pct(w["wow"])
            prev_rate = w["rate"]
        apply_comparison_fields(g)
        g["doudian_product_id"] = "、".join(sorted(g.pop("_doudian_product_ids")))
    if period == "progress":
        return sorted(groups.values(), key=lambda r: (-r["alertRank"], -(r["excessIssues"] or 0), -r["latestIssues"], r["code"]))
    return sorted(groups.values(), key=lambda r: (-r["latestIssues"], -r["total5"], r["code"]))


def aggregate_categories(rows: list[dict], weeks: list[str], order_weeks: list[int] | None = None, period: str = "closed") -> list[dict]:
    groups: dict[str, dict] = {}
    for row in rows:
        key = str(row.get("p1") or "未分类")
        g = groups.setdefault(key, {"p1": key, "latestIssues": 0, "total5": 0, "weeks": [{"label": w, "issues": 0, "orders": 0, "rate": 0.0, "wow": None, "wowText": ""} for w in weeks]})
        g["latestIssues"] += int(row.get("latestIssues") or 0)
        g["total5"] += int(row.get("total5") or 0)
        for i, w in enumerate(row.get("weeks") or []):
            g["weeks"][i]["issues"] += int(w.get("issues") or 0)
            if order_weeks and i < len(order_weeks):
                g["weeks"][i]["orders"] = int(order_weeks[i] or 0)
            else:
                g["weeks"][i]["orders"] = max(g["weeks"][i]["orders"], int(w.get("orders") or 0))
    for g in groups.values():
        prev_rate = None
        for w in g["weeks"]:
            w["rate"] = w["issues"] / w["orders"] * 100 if w["orders"] else None
            w["rateText"] = pct(w["rate"]) if w["orders"] else ""
            if w["rate"] is None or prev_rate is None:
                w["wow"] = None
                w["wowText"] = ""
            elif prev_rate == 0 and w["rate"] > 0:
                w["wow"] = 9999.0
                w["wowText"] = "上期为0"
            elif prev_rate == 0:
                w["wow"] = 0.0
                w["wowText"] = "0.00%"
            else:
                w["wow"] = (w["rate"] - prev_rate) / prev_rate * 100
                w["wowText"] = pct(w["wow"])
            prev_rate = w["rate"]
        apply_comparison_fields(g)
    if period == "progress":
        return sorted(groups.values(), key=lambda r: (-r["alertRank"], -(r["excessIssues"] or 0), -r["latestIssues"], r["p1"]))
    return sorted(groups.values(), key=lambda r: (-r["latestIssues"], r["p1"]))


XLSX_VIEW_CONFIG = {
    "category": {
        "sheet": "一级类别汇总",
        "columns": [
            ("一级类别", "p1", "text"),
            ("最新问题数", "latestIssues", "number"),
            ("最新售后率", "latestRate", "percent"),
            ("最新环比", "latestWow", "wow"),
        ],
    },
    "detail": {
        "sheet": "售后明细",
        "columns": [
            ("商家编码", "code", "text"),
            ("抖店商品ID", "doudian_product_id", "text"),
            ("商品", "name", "text"),
            ("一级", "p1", "text"),
            ("二级", "p2", "text"),
            ("问题大类", "category", "text"),
            ("最新问题数", "latestIssues", "number"),
            ("最新售后率", "latestRate", "percent"),
            ("最新环比", "latestWow", "wow"),
        ],
    },
    "warehouse": {
        "sheet": "仓库视图",
        "columns": [
            ("仓库", "warehouse", "text"),
            ("最新完整周库房问题数", "latestIssues", "number"),
            ("5周库房问题", "total5", "number"),
            ("最新完整周库房问题率", "latestRate", "percent"),
            ("库房问题率环比", "latestWow", "wow"),
        ],
    },
    "products": {
        "sheet": "商品视图",
        "columns": [
            ("商家编码", "code", "text"),
            ("抖店商品ID", "doudian_product_id", "text"),
            ("商品", "name", "text"),
            ("本周问题数", "latestIssues", "number"),
            ("5周合计", "total5", "number"),
            ("本周总售后率", "latestRate", "percent"),
            ("本周环比", "latestWow", "wow"),
        ],
    },
}

CSV_VIEW_HEADERS = {
    "category": ["p1", "latestIssues", "latestRate", "latestWowText"],
    "detail": ["code", "name", "p1", "p2", "category", "latestIssues", "latestRate", "latestWowText"],
    "warehouse": ["warehouse", "latestIssues", "total5", "latestRate", "latestWowText"],
    "products": ["code", "name", "latestIssues", "total5"],
}

PROGRESS_IDENTITY_COLUMNS = {
    "category": [("一级类别", "p1", "text")],
    "detail": [
        ("商家编码", "code", "text"), ("抖店商品ID", "doudian_product_id", "text"), ("商品", "name", "text"),
        ("一级", "p1", "text"), ("二级", "p2", "text"),
        ("问题大类", "category", "text"),
    ],
    "warehouse": [("仓库", "warehouse", "text")],
    "products": [("商家编码", "code", "text"), ("抖店商品ID", "doudian_product_id", "text"), ("商品", "name", "text")],
}

PROGRESS_METRIC_COLUMNS = [
    ("上周同期问题数", "previousIssues", "number"),
    ("上周同期订单数", "previousOrders", "number"),
    ("上周同期早期售后率", "previousRate", "percent"),
    ("本周同期问题数", "currentIssues", "number"),
    ("本周同期订单数", "currentOrders", "number"),
    ("本周同期早期售后率", "currentRate", "percent"),
    ("变化百分点", "rateDeltaPp", "pp"),
    ("相对变化", "relativeChange", "relative"),
    ("预计问题数", "expectedIssues", "decimal"),
    ("超预期问题数", "excessIssues", "decimal"),
    ("预警等级", "alertLevel", "text"),
]

WAREHOUSE_PROGRESS_METRIC_COLUMNS = [
    ("上周同期库房问题数", "previousIssues", "number"),
    ("上周同期订单数", "previousOrders", "number"),
    ("上周同期库房问题率", "previousRate", "percent"),
    ("本周同期库房问题数", "currentIssues", "number"),
    ("本周同期订单数", "currentOrders", "number"),
    ("本周同期库房问题率", "currentRate", "percent"),
    ("变化百分点", "rateDeltaPp", "pp"),
    ("相对变化", "relativeChange", "relative"),
    ("预计库房问题数", "expectedIssues", "decimal"),
    ("超预期库房问题数", "excessIssues", "decimal"),
    ("预警等级", "alertLevel", "text"),
]

PROGRESS_CSV_HEADERS = {
    view: [field for _, field, _ in columns + PROGRESS_METRIC_COLUMNS]
    for view, columns in PROGRESS_IDENTITY_COLUMNS.items()
}


def export_rows_for_view(data: dict, view: str, query: dict) -> list[dict]:
    rows = data.get("rows") or []
    weeks = data.get("weeks") or []
    period = normalize_period(data.get("periodMode") or query)
    if view == "category":
        return aggregate_categories(filtered_rows(rows, query), weeks, data.get("orderWeeks") or [], period)
    if view == "detail":
        return filtered_rows(rows, query)
    if view == "warehouse":
        q = (query.get("q") or [""])[0].strip()
        return [row for row in data.get("warehouseRows") or [] if row_matches(row, q)]
    return aggregate_products(filtered_rows(rows, query), weeks, period)


def safe_excel_text(value) -> str:
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def excel_value(row: dict, field: str, kind: str):
    value = row.get(field)
    if kind == "text":
        return safe_excel_text(value), None
    if kind == "number":
        if value in (None, ""):
            return 0, None
        return int(value) if float(value).is_integer() else float(value), None
    if kind == "percent":
        if value in (None, ""):
            return None, "0.00%"
        return float(value) / 100, "0.00%"
    if kind == "decimal":
        return (None if value in (None, "") else float(value)), "0.00"
    if kind == "pp":
        return (None if value in (None, "") else float(value)), '0.00 "个百分点"'
    if kind == "relative":
        if value is None:
            return None, "0.00%"
        if float(value) >= 9999:
            return "上期为0", None
        return float(value) / 100, "0.00%"
    if kind == "wow":
        text = str(row.get("latestWowText") or "")
        if value is None:
            return safe_excel_text(text), None
        if float(value) >= 9999:
            return safe_excel_text(text or "上期为0"), None
        return float(value) / 100, "0.00%"
    return safe_excel_text(value), None


def weekly_excel_value(week: dict, kind: str):
    if kind == "issues":
        return int(week.get("issues") or 0), None
    if kind == "orders":
        return int(week.get("orders") or 0), None
    if kind == "rate":
        value = week.get("rate")
        return (None if value in (None, "") else float(value) / 100), "0.00%"
    value = week.get("wow")
    text = str(week.get("wowText") or "")
    if value is None:
        return safe_excel_text(text), None
    if float(value) >= 9999:
        return safe_excel_text(text or "上期为0"), None
    return float(value) / 100, "0.00%"


def display_width(value) -> int:
    return sum(2 if ord(char) > 255 else 1 for char in str(value or ""))


def xlsx_bytes(view: str, rows: list[dict], weeks: list[str], period: str = "closed") -> bytes:
    period = normalize_period(period)
    config = XLSX_VIEW_CONFIG[view]
    workbook = Workbook()
    workbook.properties.creator = "售后 CRM"
    workbook.properties.title = config["sheet"]
    sheet = workbook.active
    sheet.title = config["sheet"] if period == "closed" else f"{config['sheet']}同进度"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"

    if period == "progress":
        metric_columns = WAREHOUSE_PROGRESS_METRIC_COLUMNS if view == "warehouse" else PROGRESS_METRIC_COLUMNS
        columns = list(PROGRESS_IDENTITY_COLUMNS[view]) + list(metric_columns)
    else:
        columns = list(config["columns"])
        issue_label = "库房问题数" if view == "warehouse" else "问题数"
        rate_label = "库房问题率" if view == "warehouse" else "售后率"
        for week in reversed(weeks):
            columns.extend(
                [
                    (f"{week} {issue_label}", "issues", "week"),
                    (f"{week} 订单数", "orders", "week"),
                    (f"{week} {rate_label}", "rate", "week"),
                    (f"{week} 环比", "wow", "week"),
                ]
            )

    header_fill = PatternFill("solid", fgColor="164735")
    header_font = Font(color="FFFFFF", bold=True)
    header_border = Border(bottom=Side(style="thin", color="A9BDB4"))
    body_border = Border(bottom=Side(style="hair", color="DCE5DF"))
    for column_index, (label, _, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
    sheet.row_dimensions[1].height = 30

    for row_index, row in enumerate(rows, start=2):
        output_column = 1
        base_columns = columns if period == "progress" else config["columns"]
        for _, field, kind in base_columns:
            value, number_format = excel_value(row, field, kind)
            cell = sheet.cell(row=row_index, column=output_column, value=value)
            if number_format:
                cell.number_format = number_format
            cell.alignment = Alignment(vertical="top", wrap_text=kind == "text")
            cell.border = body_border
            output_column += 1
        if period == "closed":
            row_weeks = row.get("weeks") or []
            for week_index in reversed(range(len(weeks))):
                week = row_weeks[week_index] if week_index < len(row_weeks) else {}
                for kind in ("issues", "orders", "rate", "wow"):
                    value, number_format = weekly_excel_value(week, kind)
                    cell = sheet.cell(row=row_index, column=output_column, value=value)
                    if number_format:
                        cell.number_format = number_format
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                    cell.border = body_border
                    output_column += 1

    last_column = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, len(rows) + 1)}"
    sheet.print_area = f"A1:{last_column}{max(1, len(rows) + 1)}"
    for column_index in range(1, len(columns) + 1):
        values = [sheet.cell(row=row_index, column=column_index).value for row_index in range(1, min(len(rows) + 1, 200) + 1)]
        width = max(display_width(value) for value in values) + 2
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(width, 10), 36)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class CRMHandler(BaseHTTPRequestHandler):
    server_version = "EcomCRM/0.1"

    def do_GET(self):
        self.route()

    def do_POST(self):
        self.route()

    def route(self):
        parsed = urlparse(self.path)
        query = parse_qs(parsed.query)
        path = parsed.path.rstrip("/") or "/"
        if path == "/crm/login":
            return self.login()
        if path == "/crm/logout":
            return self.logout()
        user = self.current_user()
        if not user:
            return self.redirect(f"/crm/login?next={quote(self.path)}")
        if path == "/crm":
            return self.home(user)
        if path == "/crm/dashboard":
            return self.dashboard(user, query)
        if path == "/crm/dashboard/full":
            return self.full_dashboard(user)
        if path == "/crm/dashboard/detail":
            return self.detail(user, query)
        if path == "/crm/warehouse":
            return self.warehouse(user, query)
        if path == "/crm/products":
            return self.products(user, query)
        if path == "/crm/export":
            return self.export(user, query)
        if path == "/crm/admin/users":
            return self.admin_users(user)
        self.send_error(HTTPStatus.NOT_FOUND, "Not found")

    def read_form(self) -> dict:
        length = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(length).decode("utf-8") if length else ""
        return {k: v[0] for k, v in parse_qs(raw).items()}

    def current_user(self) -> dict | None:
        cookie = SimpleCookie(self.headers.get("Cookie", ""))
        token = cookie.get("crm_session")
        if token:
            sessions = read_json(SESSION_FILE, {})
            sess = sessions.get(token.value)
            if sess and sess.get("expires_at", 0) >= time.time():
                users = read_json(USER_FILE, {})
                user = users.get(sess.get("username"))
                if user and user.get("is_active"):
                    return user
        if self.is_intranet_request():
            return INTRANET_VIEWER
        return None

    def client_ip(self) -> str:
        peer_text = str(self.client_address[0])
        try:
            peer = ipaddress.ip_address(peer_text)
        except ValueError:
            return peer_text
        if any(peer in network for network in TRUSTED_PROXY_NETWORKS):
            forwarded = (self.headers.get("X-Forwarded-For") or "").split(",", 1)[0].strip()
            if forwarded:
                try:
                    return str(ipaddress.ip_address(forwarded))
                except ValueError:
                    pass
        return peer_text

    def is_intranet_request(self) -> bool:
        try:
            ip = ipaddress.ip_address(self.client_ip())
        except ValueError:
            return False
        return any(ip in network for network in INTRANET_NETWORKS)

    def login(self):
        query = parse_qs(urlparse(self.path).query)
        if self.command == "GET" and self.is_intranet_request() and (query.get("force") or [""])[0] != "1":
            return self.redirect("/crm")
        if self.command == "POST":
            form = self.read_form()
            users = read_json(USER_FILE, {})
            user = users.get(form.get("username", "").strip())
            if user and user.get("is_active") and verify_password(form.get("password", ""), user.get("password", "")):
                token = secrets.token_urlsafe(32)
                mutate_json(
                    SESSION_FILE,
                    {},
                    lambda sessions: sessions.__setitem__(
                        token, {"username": user["username"], "expires_at": time.time() + SESSION_TTL}
                    ),
                )

                def update_last_login(current_users):
                    current = current_users.get(user["username"])
                    if current:
                        current["last_login_at"] = now_text()

                mutate_json(USER_FILE, {}, update_last_login)
                next_url = form.get("next") or "/crm"
                self.send_response(302)
                self.send_header("Location", next_url)
                self.send_header("Set-Cookie", f"crm_session={token}; Path=/crm; HttpOnly; SameSite=Lax; Max-Age={SESSION_TTL}")
                self.end_headers()
                return
            return self.html(layout("登录", login_form("用户名或密码错误", form.get("next") or "/crm"), None), HTTPStatus.UNAUTHORIZED)
        return self.html(layout("登录", login_form("", (query.get("next") or ["/crm"])[0]), None))

    def logout(self):
        cookie = SimpleCookie(self.headers.get("Cookie", ""))
        token = cookie.get("crm_session")
        if token:
            mutate_json(SESSION_FILE, {}, lambda sessions: sessions.pop(token.value, None))
        self.send_response(302)
        self.send_header("Location", "/crm/login")
        self.send_header("Set-Cookie", "crm_session=; Path=/crm; Max-Age=0")
        self.end_headers()

    def home(self, user):
        data, path = load_data()
        k = data.get("kpis") or {}
        source = source_info(data)
        body = f"""
        <section class="grid">
          {card("最新完整周", k.get("latest") or (data.get("weeks") or [""])[-1] if data.get("weeks") else "")}
          {card("趋势条目", fmt_num(k.get("rowCount") or len(data.get("rows") or [])))}
          {card("最新完整周问题数", fmt_num(k.get("totalLatest")))}
          {card("最新周环比上升条目", fmt_num(k.get("deteriorated")))}
        </section>
        <section class="card">
          <h2>导航</h2>
          <p><a class="btn" href="/crm/dashboard">日更售后 Dashboard</a> <a class="btn" href="/crm/dashboard/detail">售后明细</a> <a class="btn" href="/crm/warehouse">仓库问题视图</a> <a class="btn" href="/crm/products">商品问题视图</a></p>
          <p class="muted source">当前数据源：{esc(path)}</p>
          <p class="muted">数据周期：{esc(source.get("currentPeriod", ""))}；订单口径：{esc(source.get("orderKey", ""))}；售后去重口径：{esc(source.get("aftersalesKey", ""))}</p>
        </section>
        """
        self.html(layout("CRM 首页 / 导航", body, user))

    def dashboard(self, user, query):
        period = normalize_period(query)
        data, path = load_data(period)
        rows = data.get("rows") or []
        weeks = data.get("weeks") or []
        k = data.get("kpis") or {}
        source = source_info(data)
        category_rows = aggregate_categories(rows, weeks, data.get("orderWeeks") or [], period)
        top_rows = rows[:30]
        warehouse_rows = data.get("warehouseRows") or []
        tabs = period_tabs("/crm/dashboard", query)
        if period == "progress":
            meta = data.get("periodMeta") or {}
            completed_days = int(meta.get("completedDays") or 0)
            current_rate = pct(k.get("currentRate")) or "分母不可用"
            if completed_days:
                period_note = f"统计至昨日，共 {completed_days} 个完整自然日；今天数据未计入。"
            else:
                period_note = "本周尚无完整自然日，暂不计算同进度环比和预警。"
            body = f"""
            {tabs}
            <section class="grid">
              {card("本周已完成天数", completed_days)}
              {card("本周同期问题数", fmt_num(k.get("totalLatest")))}
              {card("本周同期早期售后率", current_rate)}
              {card("红色预警条目", fmt_num(k.get("redAlerts")))}
            </section>
            <section class="card">
              <div>{esc(period_note)}</div>
              <div class="muted source">当前数据源：{esc(path)}；数据同步：{esc(source.get("syncedAt", ""))}</div>
              <div class="muted">本周与上周均只计算相同已完成自然日，并用 api_created_at 限制到相同观察截止时间。创建时间缺失 {fmt_num(meta.get("missingCreatedAt"))} 条，已排除。</div>
            </section>
            <div class="section-title"><h2>一级类别汇总</h2><div class="section-actions">{excel_export_link("category", user, query)}</div></div>
            {progress_table(category_rows, ["p1"], {"p1": "一级类别"})}
            <div class="section-title"><h2>重点问题明细</h2><div class="section-actions"><a class="btn" href="/crm/dashboard/detail?{urlencode({'period': 'progress', 'rise': '1'})}">只看当前同进度上升</a>{excel_export_link("detail", user, query)}</div></div>
            {progress_table(top_rows, ["code", "doudian_product_id", "name", "p1", "p2", "category"], {"code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品", "p1": "一级", "p2": "二级", "category": "问题大类"})}
            <div class="section-title"><h2>仓库概览</h2><div class="section-actions"><a class="btn" href="/crm/warehouse?period=progress">查看全部仓库</a>{excel_export_link("warehouse", user, query)}</div></div>
            {progress_table(warehouse_rows, ["warehouse"], {"warehouse": "仓库"}, "库房问题")}
            """
            return self.html(layout("本周同进度 · 售后预警", body, user, period))
        body = f"""
        {tabs}
        <section class="grid">
          {card("最新完整周", k.get("latest") or (weeks[-1] if weeks else ""))}
          {card("最新完整周问题数", fmt_num(k.get("totalLatest")))}
          {card("最新周环比上升条目", fmt_num(k.get("deteriorated")))}
          {card("上期为0条目", fmt_num(k.get("zeroPrev")))}
        </section>
        <section class="card">
          <div class="muted source">当前数据源：{esc(path)}</div>
          <div class="muted">数据周期：{esc(source.get("currentPeriod", ""))}；对比窗口：近 5 个完整周；最新周售后数据仍可能继续回补。颜色：红=当期售后率 >= 2%，黄=环比售后率增长，深红=同时命中。</div>
        </section>
        <div class="section-title"><h2>一级类别汇总</h2><div class="section-actions"><span class="muted">最新完整周在左</span>{excel_export_link("category", user, query)}</div></div>
        {trend_table(category_rows, ["p1", "latestIssues", "latestRate", "latestWowText"], {"p1": "一级类别", "latestIssues": "最新完整周问题数", "latestRate": "最新完整周售后率", "latestWowText": "完整周环比"}, weeks)}
        <div class="section-title"><h2>重点问题明细</h2><div class="section-actions"><a class="btn" href="/crm/dashboard/detail?rise=1">只看 5 周内环比上升</a>{excel_export_link("detail", user, query)}</div></div>
        {trend_table(top_rows, ["code", "doudian_product_id", "name", "p1", "p2", "category", "latestIssues", "latestRate", "latestWowText"], {"code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品", "p1": "一级", "p2": "二级", "category": "问题大类", "latestIssues": "最新完整周问题数", "latestRate": "最新完整周售后率", "latestWowText": "完整周环比"}, weeks)}
        <div class="section-title"><h2>仓库概览</h2><div class="section-actions"><a class="btn" href="/crm/warehouse">查看全部仓库</a>{excel_export_link("warehouse", user, query)}</div></div>
        {trend_table(warehouse_rows, ["warehouse", "latestIssues", "latestRate", "latestWowText"], {"warehouse": "仓库", "latestIssues": "最新完整周库房问题数", "latestRate": "最新完整周库房问题率", "latestWowText": "库房问题率环比"}, weeks, "库房问题")}
        """
        self.html(layout("完整周趋势 · 售后 Dashboard", body, user, period))

    def full_dashboard(self, user):
        path = latest_dashboard_path()
        if not path:
            return self.html(layout("日更售后 Dashboard", '<div class="card">暂无售后 Dashboard 文件。</div>', user))
        content = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def detail(self, user, query):
        period = normalize_period(query)
        data, _ = load_data(period)
        rows = filtered_rows(data.get("rows") or [], query)
        weeks = data.get("weeks") or []
        body = period_tabs("/crm/dashboard/detail", query)
        body += filter_form("/crm/dashboard/detail", query, data)
        merchant_code = (query.get("merchant_code") or [""])[0].strip()
        q = (query.get("q") or [""])[0].strip()
        if merchant_code:
            body += f'<p class="ok">当前按商家编码 <b>{esc(merchant_code)}</b> 精确对账。</p>'
        elif q:
            body += '<p class="muted">当前为商品名/问题模糊检索，可能命中多个商家编码；对账请使用“商家编码（精确匹配）”。</p>'
        body += f'<div class="section-title"><h2>售后明细</h2>{excel_export_link("detail", user, query)}</div>'
        if period == "progress":
            body += progress_table(rows[:500], ["code", "doudian_product_id", "name", "p1", "p2", "category"], {
                "code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品", "p1": "一级", "p2": "二级", "category": "问题大类"
            })
        else:
            body += trend_table(rows[:500], ["code", "doudian_product_id", "name", "p1", "p2", "category", "latestIssues", "latestRate", "latestWowText"], {
                "code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品", "p1": "一级", "p2": "二级", "category": "问题大类", "latestIssues": "最新完整周问题数", "latestRate": "最新完整周售后率", "latestWowText": "完整周环比"
            }, weeks)
        body += f'<p class="muted">当前筛选 {len(rows)} 条，页面展示前 500 条。</p>'
        self.html(layout("售后明细", body, user, period))

    def warehouse(self, user, query):
        period = normalize_period(query)
        data, _ = load_data(period)
        q = (query.get("q") or [""])[0].strip()
        rows = [r for r in data.get("warehouseRows") or [] if row_matches(r, q)]
        weeks = data.get("weeks") or []
        body = period_tabs("/crm/warehouse", query)
        body += search_form("/crm/warehouse", q, period)
        body += f'<div class="section-title"><h2>仓库数据</h2>{excel_export_link("warehouse", user, query)}</div>'
        if period == "progress":
            body += progress_table(rows, ["warehouse"], {"warehouse": "仓库"}, "库房问题")
        else:
            body += trend_table(rows, ["warehouse", "latestIssues", "total5", "latestRate", "latestWowText"], {"warehouse": "仓库", "latestIssues": "最新完整周库房问题数", "total5": "5周库房问题", "latestRate": "最新完整周库房问题率", "latestWowText": "库房问题率环比"}, weeks, "库房问题")
        self.html(layout("仓库问题视图", body, user, period))

    def products(self, user, query):
        period = normalize_period(query)
        data, _ = load_data(period)
        rows = aggregate_products(filtered_rows(data.get("rows") or [], query), data.get("weeks") or [], period)
        for row in rows:
            row["_detailUrl"] = product_detail_url(row.get("code") or "", period)
        weeks = data.get("weeks") or []
        body = period_tabs("/crm/products", query)
        body += filter_form("/crm/products", query, data)
        body += '<p class="muted">点击商家编码或商品名，可按商家编码精确下钻并与商品汇总对账。</p>'
        body += f'<div class="section-title"><h2>商品数据</h2>{excel_export_link("products", user, query)}</div>'
        if period == "progress":
            body += progress_table(rows[:500], ["code", "doudian_product_id", "name"], {"code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品"})
        else:
            body += trend_table(rows[:500], ["code", "doudian_product_id", "name", "latestIssues", "total5", "latestRate", "latestWowText"], {"code": "商家编码", "doudian_product_id": "抖店商品ID", "name": "商品", "latestIssues": "最新完整周问题数", "total5": "5周合计", "latestRate": "最新完整周总售后率", "latestWowText": "完整周环比"}, weeks)
        body += f'<p class="muted">当前筛选 {len(rows)} 个商品，页面展示前 500 个。</p>'
        self.html(layout("商品问题视图", body, user, period))

    def export(self, user, query):
        view = (query.get("view") or ["detail"])[0]
        export_format = (query.get("format") or ["csv"])[0].lower()
        period = normalize_period(query)
        if view not in {"category", "detail", "warehouse", "products"} or export_format not in {"csv", "xlsx"}:
            return self.html(
                layout("导出参数错误", '<section class="card"><p class="bad">未知的导出视图或格式。</p></section>', user),
                HTTPStatus.BAD_REQUEST,
            )
        if export_format == "xlsx" and not can_export_excel(user):
            return self.html(layout("无导出权限", '<section class="card"><p class="bad">请登录后导出 Excel。</p></section>', user), HTTPStatus.FORBIDDEN)
        if export_format == "csv" and not can_export(user):
            return self.html(layout("无导出权限", '<section class="card"><p class="bad">无导出权限</p></section>', user), HTTPStatus.FORBIDDEN)

        data, _ = load_data(period)
        rows = export_rows_for_view(data, view, query)
        if export_format == "xlsx":
            content = xlsx_bytes(view, rows, data.get("weeks") or [], period)
            extension = "xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            headers = PROGRESS_CSV_HEADERS[view] if period == "progress" else CSV_VIEW_HEADERS[view]
            content = csv_bytes(rows, headers)
            extension = "csv"
            content_type = "text/csv; charset=utf-8-sig"
        period_name = "本周同进度" if period == "progress" else "完整周趋势"
        ascii_filename = f"crm_{view}_{period}_{datetime.now():%Y%m%d_%H%M%S}.{extension}"
        display_filename = f"CRM_{view}_{period_name}_{datetime.now():%Y%m%d_%H%M%S}.{extension}"
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f"attachment; filename=\"{ascii_filename}\"; filename*=UTF-8''{quote(display_filename)}")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def admin_users(self, user):
        if not is_admin(user):
            return self.html(
                layout("无管理员权限", '<section class="card"><p class="bad">无管理员权限</p></section>', user),
                HTTPStatus.FORBIDDEN,
            )
        msg = ""
        if self.command == "POST":
            form = self.read_form()
            username = form.get("username", "").strip()
            password = form.get("password", "")
            role = form.get("role", "viewer")
            if username and password and role in ROLES:
                new_user = {"username": username, "display_name": form.get("display_name", username), "role": role, "is_active": True, "password": hash_password(password), "created_at": now_text(), "updated_at": now_text(), "last_login_at": ""}
                mutate_json(USER_FILE, {}, lambda current_users: current_users.__setitem__(username, new_user))
                msg = "用户已创建"
        users = read_json(USER_FILE, {})
        body = admin_users_page(users, msg)
        self.html(layout("用户管理", body, user))

    def redirect(self, target):
        self.send_response(302)
        self.send_header("Location", target)
        self.end_headers()

    def html(self, content: bytes, status=HTTPStatus.OK):
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, fmt, *args):
        sys.stderr.write("%s %s\n" % (now_text(), fmt % args))


def card(label, value):
    return f'<div class="card"><div class="label">{esc(label)}</div><div class="value">{esc(value)}</div></div>'


def period_tabs(path: str, query: dict) -> str:
    current = normalize_period(query)
    retained = []
    for key in ("q", "merchant_code", "doudian_product_id", "p1", "p2", "rise"):
        value = (query.get(key) or [""])[0]
        if value:
            retained.append((key, value))
    links = []
    for period, label in (("closed", "完整周趋势"), ("progress", "本周同进度")):
        params = retained + [("period", period)]
        active = " active" if current == period else ""
        links.append(f'<a class="btn{active}" href="{esc(path)}?{esc(urlencode(params))}">{label}</a>')
    return '<div class="period-tabs">' + "".join(links) + "</div>"


def login_form(message: str, next_url: str) -> str:
    error = f'<p class="bad">{esc(message)}</p>' if message else ""
    return f"""<section class="card login">
      <h2>登录</h2>{error}
      <form method="post" action="/crm/login">
        <input type="hidden" name="next" value="{esc(next_url)}">
        <p><input name="username" placeholder="用户名" autofocus></p>
        <p><input name="password" placeholder="密码" type="password"></p>
        <p><button type="submit">登录</button></p>
      </form>
    </section>"""


def search_form(action: str, q: str, period: str = "closed") -> str:
    return f'<form class="toolbar" method="get" action="{esc(action)}"><input type="hidden" name="period" value="{esc(normalize_period(period))}"><input name="q" value="{esc(q)}" placeholder="搜索"><button type="submit">筛选</button></form>'


def select_options(values: list[str], selected: str, blank: str) -> str:
    opts = [f'<option value="">{esc(blank)}</option>']
    for value in values:
        sel = " selected" if value == selected else ""
        opts.append(f'<option value="{esc(value)}"{sel}>{esc(value)}</option>')
    return "".join(opts)


def filter_form(action: str, query: dict, data: dict | None = None) -> str:
    q = (query.get("q") or [""])[0]
    merchant_code = (query.get("merchant_code") or [""])[0]
    doudian_product_id = (query.get("doudian_product_id") or [""])[0]
    p1 = (query.get("p1") or [""])[0]
    p2 = (query.get("p2") or [""])[0]
    rise = (query.get("rise") or [""])[0]
    filters = (data or {}).get("filters") or {}
    p1_options = select_options(filters.get("p1") or [], p1, "全部一级")
    p2_options = select_options(filters.get("p2") or [], p2, "全部二级")
    checked = " checked" if rise == "1" else ""
    period = normalize_period(query)
    rise_label = "只看当前同进度上升" if period == "progress" else "只看5周内环比上升"
    return f"""<form class="toolbar" method="get" action="{esc(action)}">
      <input type="hidden" name="period" value="{esc(period)}">
      <input name="merchant_code" value="{esc(merchant_code)}" placeholder="商家编码（精确匹配，对账用）">
      <input name="q" value="{esc(q)}" placeholder="模糊搜索商品名 / 问题">
      <input name="doudian_product_id" value="{esc(doudian_product_id)}" placeholder="抖店商品ID（精确匹配）">
      <select name="p1">{p1_options}</select>
      <select name="p2">{p2_options}</select>
      <label class="pill"><input type="checkbox" name="rise" value="1"{checked}> {rise_label}</label>
      <button type="submit">筛选</button>
    </form>"""


def rows_table(rows: list[dict], fields: list[str], labels: dict[str, str]) -> str:
    head = "".join(f"<th>{esc(labels.get(f, f))}</th>" for f in fields)
    body = []
    for row in rows:
        cells = []
        for f in fields:
            value = row.get(f, "")
            if f.endswith("Rate") or f == "latestRate":
                value = pct(value)
            elif isinstance(value, (int, float)) and "Rate" not in f:
                value = fmt_num(value)
            cells.append(f"<td>{esc(value)}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    body_html = "".join(body) or '<tr><td colspan="20">暂无数据</td></tr>'
    return f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{body_html}</tbody></table></div>'


def display_value(row: dict, field: str) -> str:
    value = row.get(field, "")
    if field.endswith("Rate") or field == "latestRate":
        rendered = pct(value)
    elif isinstance(value, (int, float)) and "Rate" not in field:
        rendered = fmt_num(value)
    else:
        rendered = esc(value)
    detail_url = str(row.get("_detailUrl") or "")
    if detail_url and field in {"code", "name"}:
        return f'<a href="{esc(detail_url)}" title="按商家编码精确查看售后明细">{rendered}</a>'
    return rendered


def trend_table(rows: list[dict], fields: list[str], labels: dict[str, str], weeks: list[str], metric_label: str = "问题") -> str:
    week_headers = "".join(f'<th>{esc(label)}<br><span class="muted">{esc(metric_label)}/订单/率/环比</span></th>' for label in reversed(weeks))
    base_headers = "".join(f'<th class="base-col">{esc(labels.get(f, f))}</th>' for f in fields)
    body = []
    for row in rows:
        cells = []
        for f in fields:
            cls = "wide-name" if f == "name" else ""
            cells.append(f'<td class="{cls}">{display_value(row, f)}</td>')
        row_weeks = row.get("weeks") or []
        for idx in reversed(range(len(weeks))):
            week = row_weeks[idx] if idx < len(row_weeks) else {"issues": 0, "rate": None, "wow": None}
            cls = heat_class(week)
            issues = fmt_num(week.get("issues"))
            orders = fmt_num(week.get("orders"))
            rate = week_rate_text(week)
            wow = week_wow_text(week)
            cells.append(f'<td class="week-cell {cls}"><b>{issues}</b><br><span class="muted">{orders}</span><br>{esc(rate)}<br>{esc(wow)}</td>')
        body.append("<tr>" + "".join(cells) + "</tr>")
    body_html = "".join(body) or '<tr><td colspan="30">暂无数据</td></tr>'
    return f'<div class="table-wrap"><table><thead><tr>{base_headers}{week_headers}</tr></thead><tbody>{body_html}</tbody></table></div>'


def progress_table(rows: list[dict], fields: list[str], labels: dict[str, str], metric_label: str = "问题") -> str:
    metric_headers = [
        f"上周同期{metric_label}", "上周同期订单", f"上周同期{metric_label}率",
        f"本周同期{metric_label}", "本周同期订单", f"本周同期{metric_label}率",
        "变化百分点", "相对变化", f"预计{metric_label}", f"超预期{metric_label}", "预警",
    ]
    headers = "".join(f'<th class="base-col">{esc(labels.get(field, field))}</th>' for field in fields)
    headers += "".join(f"<th>{esc(label)}</th>" for label in metric_headers)
    body = []
    for row in rows:
        cells = [f'<td class="{"wide-name" if field == "name" else ""}">{display_value(row, field)}</td>' for field in fields]
        previous_rate = pct(row.get("previousRate")) or "分母不可用"
        current_rate = pct(row.get("currentRate")) or "分母不可用"
        delta = row.get("rateDeltaPp")
        delta_text = "—" if delta is None else f"{float(delta):+.2f} 个百分点"
        relative = row.get("relativeChange")
        if relative is None:
            relative_text = "—"
        elif float(relative) >= 9999:
            relative_text = "上期为0"
        else:
            relative_text = f"{float(relative):+.2f}%"
        expected = row.get("expectedIssues")
        excess = row.get("excessIssues")
        expected_text = "—" if expected is None else f"{float(expected):.2f}"
        excess_text = "—" if excess is None else f"{float(excess):+.2f}"
        level = str(row.get("alertLevel") or "灰色")
        level_class = {"红色": "alert-red", "黄色": "alert-yellow", "绿色": "alert-green"}.get(level, "alert-gray")
        values = [
            fmt_num(row.get("previousIssues")), fmt_num(row.get("previousOrders")), previous_rate,
            fmt_num(row.get("currentIssues")), fmt_num(row.get("currentOrders")), current_rate,
            delta_text, relative_text, expected_text, excess_text,
        ]
        cells.extend(f"<td>{esc(value)}</td>" for value in values)
        cells.append(f'<td class="{level_class}">{esc(level)}</td>')
        body.append("<tr>" + "".join(cells) + "</tr>")
    body_html = "".join(body) or '<tr><td colspan="30">暂无同进度数据</td></tr>'
    return f'<div class="table-wrap"><table><thead><tr>{headers}</tr></thead><tbody>{body_html}</tbody></table></div>'


def csv_bytes(rows: list[dict], headers: list[str]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(h, "") for h in headers])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def admin_users_page(users: dict, msg: str) -> str:
    rows = "".join(
        f"<tr><td>{esc(u.get('username'))}</td><td>{esc(u.get('display_name'))}</td><td>{esc(role_label(u.get('role','')))}</td><td>{'启用' if u.get('is_active') else '停用'}</td><td>{esc(u.get('last_login_at'))}</td></tr>"
        for u in sorted(users.values(), key=lambda x: x.get("username", ""))
    )
    return f"""<section class="card">
      <h2>新增用户</h2>
      {'<p class="ok">' + esc(msg) + '</p>' if msg else ''}
      <form class="toolbar" method="post">
        <input name="username" placeholder="用户名">
        <input name="display_name" placeholder="姓名">
        <input name="password" placeholder="初始密码">
        <select name="role">
          <option value="viewer">只读访问</option>
          <option value="exporter">导出权限</option>
          <option value="admin">管理员</option>
        </select>
        <button type="submit">创建</button>
      </form>
    </section>
    <section class="card"><h2>用户列表</h2><table><thead><tr><th>用户名</th><th>姓名</th><th>角色</th><th>状态</th><th>最后登录</th></tr></thead><tbody>{rows}</tbody></table></section>"""


def main():
    ensure_runtime()
    server = ThreadingHTTPServer((HOST, PORT), CRMHandler)
    print(f"CRM dashboard listening on http://{HOST}:{PORT}/crm", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()

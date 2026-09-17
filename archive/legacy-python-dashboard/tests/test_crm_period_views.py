import io
import sys
import types
import unittest
from datetime import date

from openpyxl import load_workbook

sys.modules.setdefault("pymysql", types.ModuleType("pymysql"))

import crm_mysql_data
import crm_server
from scripts.crm_order_daily_refresh import day_windows


def progress_row() -> dict:
    row = {
        "code": "P001",
        "product_id": "DOUDIAN-9988",
        "doudian_product_id": "DOUDIAN-9988",
        "name": "测试商品",
        "p1": "质量问题",
        "p2": "破损",
        "p3": "",
        "category": "质量问题",
        "weeks": [
            {"issues": 6, "orders": 800, "rate": 0.75, "rateText": "0.75%", "wow": None, "wowText": ""},
            {"issues": 12, "orders": 1000, "rate": 1.2, "rateText": "1.20%", "wow": 60.0, "wowText": "60.00%"},
        ],
        "total5": 18,
    }
    return crm_server.apply_comparison_fields(row)


class PeriodWindowTests(unittest.TestCase):
    def test_monday_compares_one_completed_day(self):
        windows = crm_mysql_data.progress_windows(date(2026, 8, 3))
        self.assertEqual((windows[1]["start"], windows[1]["end"]), (date(2026, 8, 2), date(2026, 8, 3)))
        self.assertEqual((windows[0]["start"], windows[0]["end"]), (date(2026, 7, 26), date(2026, 7, 27)))
        self.assertEqual(windows[0]["observe_end"], date(2026, 7, 27))

    def test_sunday_has_no_completed_day(self):
        windows = crm_mysql_data.progress_windows(date(2026, 8, 2))
        self.assertEqual(windows[1]["start"], windows[1]["end"])
        self.assertEqual(windows[0]["start"], windows[0]["end"])
        self.assertIn("无完整日", windows[1]["label"])

    def test_cross_year_window_and_daily_refresh_range(self):
        windows = crm_mysql_data.progress_windows(date(2026, 1, 1))
        self.assertEqual(windows[1]["start"], date(2025, 12, 28))
        days = day_windows(70, date(2026, 1, 1))
        self.assertEqual(days[0][0], date(2025, 10, 23))
        self.assertEqual(days[-1][1], date(2026, 1, 1))

    def test_progress_sql_uses_equal_observation_cutoffs(self):
        windows = crm_mysql_data.progress_windows(date(2026, 8, 5))
        sql = crm_mysql_data.weekly_case("a", windows)
        params = crm_mysql_data.window_params(windows)
        self.assertEqual(sql.count("api_created_at < %s"), 2)
        self.assertEqual(len(params), 6)
        self.assertEqual(params[2], windows[0]["observe_end"])
        self.assertEqual(params[5], windows[1]["observe_end"])

    def test_zero_denominator_is_not_rendered_as_zero_percent(self):
        weeks = crm_mysql_data.build_weeks([2, 3], [0, 0])
        self.assertIsNone(weeks[0]["rate"])
        self.assertIsNone(weeks[1]["rate"])
        self.assertEqual(weeks[1]["rateText"], "")
        fields = crm_mysql_data.latest_fields(weeks)
        self.assertIsNone(fields["rateDeltaPp"])
        self.assertEqual(fields["alertLevel"], "灰色")

    def test_platform_product_denominator_is_preserved(self):
        values = crm_mysql_data.product_order_series(
            "MERCHANT-CODE",
            {"MERCHANT-CODE": [10, 20]},
            [100, 200],
            "DOUDIAN-ID",
            {"DOUDIAN-ID": [30, 40]},
        )
        self.assertEqual(values, [30, 40])

    def test_unmatched_product_does_not_fall_back_to_all_orders(self):
        values = crm_mysql_data.product_order_series("未填", {}, [100, 200])
        self.assertEqual(values, [0, 0])

    def test_warehouse_row_uses_only_warehouse_problem_counts(self):
        row = crm_mysql_data.build_warehouse_row("QHC", "启航仓", [172, 173], [146_558, 151_262])
        self.assertEqual(row["warehouse"], "QHC / 启航仓")
        self.assertEqual(row["latestIssues"], 173)
        self.assertEqual(row["total5"], 345)
        self.assertAlmostEqual(row["latestRate"], 173 / 151_262 * 100)
        html = crm_server.trend_table(
            [row],
            ["warehouse", "latestIssues", "latestRate"],
            {"warehouse": "仓库", "latestIssues": "库房问题数", "latestRate": "库房问题率"},
            ["0705-0711", "0712-0718"],
            "库房问题",
        )
        self.assertIn("库房问题/订单/率/环比", html)
        self.assertIn("173", html)

    def test_unmapped_warehouse_rate_is_unavailable(self):
        row = crm_mysql_data.build_warehouse_row("UNKNOWN", "", [2, 3], [0, 0])
        self.assertIsNone(row["latestRate"])
        self.assertEqual(row["weeks"][-1]["rateText"], "")


class ProgressPresentationTests(unittest.TestCase):
    def test_visible_detail_tables_do_not_render_third_level_column(self):
        row = progress_row()
        row["p3"] = "三级问题"
        trend = crm_server.trend_table(
            [row],
            ["p1", "p2", "category"],
            {"p1": "一级", "p2": "二级", "category": "问题大类"},
            ["上周同期", "本周同期"],
        )
        progress = crm_server.progress_table(
            [row],
            ["p1", "p2", "category"],
            {"p1": "一级", "p2": "二级", "category": "问题大类"},
        )
        self.assertNotIn("三级", trend)
        self.assertNotIn("三级", progress)

    def test_alert_calculation_and_progress_table(self):
        row = progress_row()
        self.assertAlmostEqual(row["rateDeltaPp"], 0.45)
        self.assertAlmostEqual(row["expectedIssues"], 7.5)
        self.assertAlmostEqual(row["excessIssues"], 4.5)
        self.assertEqual(row["alertLevel"], "红色")
        html = crm_server.progress_table([row], ["code", "name"], {"code": "编码", "name": "商品"})
        self.assertIn("+0.45 个百分点", html)
        self.assertIn("alert-red", html)

    def test_progress_workbook_has_comparison_columns_and_numeric_formats(self):
        row = progress_row()
        content = crm_server.xlsx_bytes("detail", [row], ["上周同期", "本周同期"], "progress")
        sheet = load_workbook(io.BytesIO(content), data_only=False).active
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("变化百分点", headers)
        self.assertIn("超预期问题数", headers)
        self.assertIn("预警等级", headers)
        delta_cell = sheet.cell(2, headers.index("变化百分点") + 1)
        self.assertAlmostEqual(delta_cell.value, 0.45)
        self.assertIn("个百分点", delta_cell.number_format)

    def test_period_is_preserved_in_tabs_filters_and_exports(self):
        query = {"period": ["progress"], "doudian_product_id": ["DOUDIAN-9988"], "p1": ["质量问题"], "rise": ["1"]}
        tabs = crm_server.period_tabs("/crm/dashboard/detail", query)
        form = crm_server.filter_form("/crm/dashboard/detail", query, {"filters": {"p1": ["质量问题"]}})
        link = crm_server.excel_export_link("detail", {"username": "viewer", "role": "viewer"}, query)
        self.assertIn("period=progress", tabs)
        self.assertIn("doudian_product_id=DOUDIAN-9988", tabs)
        self.assertIn('name="period" value="progress"', form)
        self.assertIn('name="doudian_product_id" value="DOUDIAN-9988"', form)
        self.assertIn("period=progress", link)
        self.assertIn("doudian_product_id=DOUDIAN-9988", link)
        self.assertEqual(crm_server.normalize_period("unexpected"), "closed")


if __name__ == "__main__":
    unittest.main()

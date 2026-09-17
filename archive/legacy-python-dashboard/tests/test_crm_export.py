import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import crm_server


WEEKS = ["2026-W29", "2026-W30"]


def trend(code, name, p1, p2, *, rise=True):
    return {
        "code": code,
        "product_id": f"PID-{code}",
        "doudian_product_id": f"PID-{code}",
        "name": name,
        "p1": p1,
        "p2": p2,
        "p3": "三级",
        "category": "问题大类",
        "latestIssues": 3,
        "total5": 5,
        "latestRate": 1.5,
        "latestWow": 50.0 if rise else -10.0,
        "latestWowText": "50.00%" if rise else "-10.00%",
        "weeks": [
            {"label": WEEKS[0], "issues": 2, "orders": 200, "rate": 1.0, "wow": None, "wowText": ""},
            {"label": WEEKS[1], "issues": 3, "orders": 200, "rate": 1.5, "wow": 50.0 if rise else -10.0, "wowText": "50.00%" if rise else "-10.00%"},
        ],
    }


DATA = {
    "weeks": WEEKS,
    "orderWeeks": [400, 400],
    "rows": [
        trend("001", "=SUM(A1:A2)", "质量", "破损", rise=True),
        trend("002", "普通商品", "物流", "超时", rise=False),
    ],
    "warehouseRows": [
        {
            "warehouse": "@华东仓",
            "latestIssues": 3,
            "total5": 5,
            "latestRate": 0.75,
            "latestWow": 50.0,
            "latestWowText": "50.00%",
            "weeks": [
                {"issues": 2, "orders": 400, "rate": 0.5, "wow": None, "wowText": ""},
                {"issues": 3, "orders": 400, "rate": 0.75, "wow": 50.0, "wowText": "50.00%"},
            ],
        }
    ],
}


class ExportWorkbookTests(unittest.TestCase):
    def workbook(self, view, rows):
        return load_workbook(io.BytesIO(crm_server.xlsx_bytes(view, rows, WEEKS)), data_only=False)

    def test_all_four_views_have_expected_row_counts_and_workbook_features(self):
        expected = {"category": 2, "detail": 2, "warehouse": 1, "products": 2}
        for view, count in expected.items():
            with self.subTest(view=view):
                rows = crm_server.export_rows_for_view(DATA, view, {})
                workbook = self.workbook(view, rows)
                sheet = workbook.active
                self.assertEqual(sheet.max_row, count + 1)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertTrue(sheet.auto_filter.ref.startswith("A1:"))
                self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "164735")

    def test_detail_types_percent_format_week_order_and_formula_protection(self):
        rows = crm_server.export_rows_for_view(DATA, "detail", {})
        sheet = self.workbook("detail", rows).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet["A1"].value, "商家编码")
        self.assertIn("2026-W30 问题数", headers)
        self.assertEqual(sheet["C2"].value, "'=SUM(A1:A2)")
        latest_issues = sheet.cell(2, headers.index("最新问题数") + 1)
        latest_rate = sheet.cell(2, headers.index("最新售后率") + 1)
        week_issues = sheet.cell(2, headers.index("2026-W30 问题数") + 1)
        week_rate = sheet.cell(2, headers.index("2026-W30 售后率") + 1)
        self.assertIsInstance(latest_issues.value, int)
        self.assertAlmostEqual(latest_rate.value, 0.015)
        self.assertEqual(latest_rate.number_format, "0.00%")
        self.assertIsInstance(week_issues.value, int)
        self.assertAlmostEqual(week_rate.value, 0.015)
        self.assertEqual(week_rate.number_format, "0.00%")

    def test_formula_injection_prefixes_are_neutralized(self):
        for value in ("=1+1", "+1", "-1", "@cmd", "  =1+1"):
            with self.subTest(value=value):
                self.assertTrue(crm_server.safe_excel_text(value).startswith("'"))
        self.assertEqual(crm_server.safe_excel_text("正常文本"), "正常文本")

    def test_warehouse_export_uses_warehouse_problem_metrics(self):
        rows = crm_server.export_rows_for_view(DATA, "warehouse", {})
        sheet = self.workbook("warehouse", rows).active
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("最新完整周库房问题数", headers)
        self.assertIn("2026-W30 库房问题数", headers)
        self.assertIn("2026-W30 库房问题率", headers)
        self.assertNotIn("本周总售后问题数", headers)
        self.assertEqual(sheet["B2"].value, 3)

    def test_detail_exports_do_not_include_third_level_problem(self):
        rows = crm_server.export_rows_for_view(DATA, "detail", {})
        sheet = self.workbook("detail", rows).active
        headers = [cell.value for cell in sheet[1]]
        self.assertNotIn("三级", headers)
        self.assertNotIn("p3", crm_server.CSV_VIEW_HEADERS["detail"])
        self.assertNotIn("p3", [field for _, field, _ in crm_server.PROGRESS_IDENTITY_COLUMNS["detail"]])

    def test_filters_and_full_export_beyond_page_limit(self):
        many = [trend(f"{index:04d}", f"商品{index}", "质量", "破损", rise=index % 2 == 0) for index in range(620)]
        data = {"weeks": WEEKS, "orderWeeks": [1000, 1000], "rows": many, "warehouseRows": DATA["warehouseRows"]}
        filtered = crm_server.export_rows_for_view(data, "detail", {"p1": ["质量"], "p2": ["破损"], "rise": ["1"]})
        self.assertEqual(len(filtered), 310)
        all_rows = crm_server.export_rows_for_view(data, "detail", {})
        self.assertEqual(len(all_rows), 620)
        self.assertEqual(self.workbook("detail", all_rows).active.max_row, 621)
        self.assertEqual(len(crm_server.export_rows_for_view(data, "products", {"q": ["商品61"]})), 11)
        self.assertEqual(len(crm_server.export_rows_for_view(data, "detail", {"doudian_product_id": ["PID-0061"]})), 1)
        self.assertEqual(len(crm_server.export_rows_for_view(data, "detail", {"doudian_product_id": ["0061"]})), 0)
        self.assertEqual(len(crm_server.export_rows_for_view(data, "warehouse", {"q": ["华东"]})), 1)

    def test_json_writes_and_lock_files_are_private(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "users.json"
            crm_server.write_json(target, {"a": 1})
            crm_server.mutate_json(target, {}, lambda payload: payload.__setitem__("b", 2))
            self.assertEqual(crm_server.read_json(target, {}), {"a": 1, "b": 2})
            self.assertEqual(os.stat(target).st_mode & 0o777, 0o600)
            self.assertEqual(os.stat(Path(str(target) + ".lock")).st_mode & 0o777, 0o600)


class FakeExportHandler:
    export = crm_server.CRMHandler.export

    def __init__(self):
        self.status = None
        self.headers_sent = {}
        self.wfile = io.BytesIO()

    def html(self, content, status=200):
        self.status = int(status)
        self.wfile.write(content)

    def send_response(self, status):
        self.status = int(status)

    def send_header(self, name, value):
        self.headers_sent[name] = value

    def end_headers(self):
        pass


class HttpPermissionTests(unittest.TestCase):
    def export_status(self, user, query):
        handler = FakeExportHandler()
        with patch.object(crm_server, "load_data", return_value=(DATA, Path("fixture"))):
            handler.export(user, query)
        return handler

    def test_excel_role_matrix_and_legacy_csv(self):
        for role in ("viewer", "exporter", "admin"):
            with self.subTest(role=role):
                handler = self.export_status({"username": role, "role": role}, {"view": ["detail"], "format": ["xlsx"]})
                self.assertEqual(handler.status, 200)
                self.assertEqual(
                    handler.headers_sent["Content-Type"],
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
        intranet = self.export_status(crm_server.INTRANET_VIEWER, {"view": ["detail"], "format": ["xlsx"]})
        self.assertEqual(intranet.status, 403)
        viewer_csv = self.export_status({"username": "viewer", "role": "viewer"}, {"view": ["detail"]})
        self.assertEqual(viewer_csv.status, 403)
        exporter_csv = self.export_status({"username": "exporter", "role": "exporter"}, {"view": ["detail"]})
        self.assertEqual(exporter_csv.status, 200)
        with patch.object(crm_server, "XLSX_EXPORT_ENABLED", False):
            stable = self.export_status({"username": "admin", "role": "admin"}, {"view": ["detail"], "format": ["xlsx"]})
            self.assertEqual(stable.status, 403)

    def test_unknown_view_and_format_are_400(self):
        user = {"username": "admin", "role": "admin"}
        self.assertEqual(self.export_status(user, {"view": ["unknown"], "format": ["xlsx"]}).status, 400)
        self.assertEqual(self.export_status(user, {"view": ["detail"], "format": ["pdf"]}).status, 400)

    def test_proxy_header_is_only_trusted_from_loopback(self):
        direct = type("Client", (), {"headers": {"X-Forwarded-For": "8.8.8.8"}, "client_address": ("192.168.3.55", 1234)})()
        proxied = type("Client", (), {"headers": {"X-Forwarded-For": "8.8.8.8"}, "client_address": ("127.0.0.1", 1234)})()
        self.assertEqual(crm_server.CRMHandler.client_ip(direct), "192.168.3.55")
        self.assertEqual(crm_server.CRMHandler.client_ip(proxied), "8.8.8.8")

    def test_unauthenticated_external_request_redirects_to_login(self):
        class ExternalRequest:
            route = crm_server.CRMHandler.route
            current_user = crm_server.CRMHandler.current_user
            client_ip = crm_server.CRMHandler.client_ip
            is_intranet_request = crm_server.CRMHandler.is_intranet_request

            path = "/crm"
            headers = {}
            client_address = ("8.8.8.8", 1234)

            def redirect(self, target):
                self.redirected_to = target

        request = ExternalRequest()
        request.route()
        self.assertEqual(request.redirected_to, "/crm/login?next=/crm")


if __name__ == "__main__":
    unittest.main()

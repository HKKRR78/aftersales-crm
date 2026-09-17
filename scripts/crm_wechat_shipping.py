"""Recover exact SKU shipment evidence from retained, audited native exports.

The company loader intentionally clears raw_order_json. Read the protected
original without putting buyer/address fields back into the normalized DB.
"""
import hashlib,json
from collections import defaultdict
from datetime import datetime,timedelta,timezone
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from crm_report_model import text,canonical
from crm_native_snapshot import source_datetime

SHANGHAI=timezone(timedelta(hours=8))


def original_catalog(conn,wanted):
    from crm_sales_sources import query
    catalog=defaultdict(list)
    for record in query(conn,"SELECT platform_shop_id,manifest_path,manifest_sha256 FROM audit_collection_run WHERE source_code='wechat_order_api' AND release_status='READY' ORDER BY completed_at DESC"):
        path=Path(record['manifest_path'])
        if not path.is_file():continue
        raw=path.read_bytes();manifest=json.loads(raw)
        for artifact in manifest.get('artifacts',[]):
            original=Path(artifact['path']);key=(record['platform_shop_id'],original.name)
            if artifact.get('artifact_type')!='raw_export' or key not in wanted:continue
            if hashlib.sha256(raw).hexdigest()!=record['manifest_sha256']:
                raise ValueError('WeChat collection manifest changed: '+str(path))
            shop=manifest.get('shop',{})
            if not shop.get('identity_verified') or text(shop.get('platform_shop_id'))!=record['platform_shop_id']:
                raise ValueError('WeChat original shop identity is not verified')
            proof={'file':str(original),'sha256':artifact['sha256'],'manifest':str(path),'manifest_sha256':record['manifest_sha256']}
            if proof not in catalog[key]:catalog[key].append(proof)
    # Earlier company exports retain their per-file API evidence even when a
    # multi-shard run's central manifest points only to its final shard.
    # Registered files certify exact order versions, never payment coverage.
    from crm_wechat_artifact import registered_artifacts
    for key,proof in registered_artifacts(conn,wanted):
        if proof not in catalog[key]:catalog[key].append(proof)
    return catalog


def from_original(row):
    """Keep only exact original identity, purchase facts and shipment evidence."""
    raw=json.loads(row['raw_order_json']);detail=raw.get('order_detail') or {}
    key=tuple(text(row.get(k)) for k in ('order_id','product_id','sku_id'))
    if text(raw.get('order_id'))!=key[0]:raise ValueError('WeChat original order identity conflict')
    products=[p for p in detail.get('product_infos',[]) if (text(p.get('product_id')),text(p.get('sku_id')))==key[1:]]
    if len(products)!=1:raise ValueError('WeChat original SKU is absent or duplicated')
    product=products[0];qty=Decimal(str(row['sku_cnt']));code=text(row.get('sku_code'))
    if text(product.get('sku_code'))!=code or Decimal(str(product.get('sku_cnt')))!=qty:
        raise ValueError('WeChat original purchase facts differ from exported columns')
    def timestamp(value):
        return datetime.fromtimestamp(int(value),SHANGHAI).replace(tzinfo=None) if value else None
    paid=source_datetime(row.get('pay_time'));updated=source_datetime(row.get('update_time'))
    if paid!=timestamp((detail.get('pay_info') or {}).get('pay_time')) or updated!=timestamp(raw.get('update_time')):
        raise ValueError('WeChat original timestamps differ from exported columns')
    deliveries=[]
    for shipment in (detail.get('delivery_info') or {}).get('delivery_product_info',[]):
        deliveries.append({k:shipment.get(k) for k in ('delivery_id','waybill_id','delivery_time','product_infos')})
    return key,{'code':code,'quantity':qty,'paid_at':paid,'version_at':updated,'shipments':deliveries,
                'original_hash':hashlib.sha256(canonical(raw).encode()).hexdigest()}


def original_rows(filename,digest):
    stat=Path(filename).stat()
    return read_original_rows(filename,digest,(stat.st_ino,stat.st_size,stat.st_mtime_ns))


@lru_cache(maxsize=32)
def read_original_rows(filename,digest,fingerprint):
    from openpyxl import load_workbook
    path=Path(filename)
    if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('WeChat original export hash changed: '+filename)
    before=path.stat();workbook=load_workbook(path,read_only=True,data_only=True);records={}
    if (before.st_ino,before.st_size,before.st_mtime_ns)!=fingerprint:raise ValueError('WeChat original changed before read')
    try:
        rows=workbook['订单'].iter_rows(values_only=True);headers=next(rows)
        required={'order_id','product_id','sku_id','sku_code','sku_cnt','pay_time','update_time','raw_order_json'}
        if not required.issubset(headers):raise ValueError('WeChat original business columns missing')
        for ordinal,values in enumerate(rows,2):
            row=dict(zip(headers,values))
            if not row.get('order_id'):continue
            if not row.get('raw_order_json'):continue
            key,record=from_original(row)
            if key in records:raise ValueError('duplicate WeChat original sales line')
            records[key]={**record,'row':ordinal}
    finally:workbook.close()
    after=path.stat()
    if (before.st_ino,before.st_size,before.st_mtime_ns)!=(after.st_ino,after.st_size,after.st_mtime_ns):
        raise ValueError('WeChat original changed during read')
    return records


def restore_shipments(conn,rows):
    targets=defaultdict(list)
    for row in rows:
        if row.get('native_shipments') is None:targets[(text(row['shop_id']),Path(row['source_file']).name)].append(row)
    if not targets:return
    catalog=original_catalog(conn,set(targets))
    for file_key,wanted in targets.items():
        candidates=defaultdict(list)
        reasons=defaultdict(set)
        existing_files=0
        for proof in catalog.get(file_key,[]):
            if not Path(proof['file']).is_file():continue
            existing_files+=1
            records=original_rows(proof['file'],proof['sha256'])
            for row in wanted:
                key=tuple(text(row.get(k)) for k in ('order_no','product_id','native_sku_id'))
                original=records.get(key)
                if original is None:
                    reasons[key].add('native_original_sales_line_absent');continue
                if (original['code']==text(row['code']) and original['quantity']==Decimal(str(row['quantity']))
                    and original['paid_at']==source_datetime(row['paid_at']) and original['version_at']==source_datetime(row['version_at'])):
                    candidates[key].append((original,{**proof,'row':original['row']}))
                else:reasons[key].add('native_original_exact_version_mismatch')
        for row in wanted:
            key=tuple(text(row.get(k)) for k in ('order_no','product_id','native_sku_id'));found=candidates[key]
            if not found:
                row['shipping_source_reason']=(
                    'native_original_catalog_missing' if file_key not in catalog else
                    'native_original_file_unavailable' if not existing_files else
                    'native_original_exact_version_mismatch' if 'native_original_exact_version_mismatch' in reasons[key] else
                    'native_original_sales_line_absent')
                continue
            if len({canonical(item['shipments']) for item,_ in found})!=1:
                raise ValueError('conflicting WeChat original shipment evidence for the same order version')
            original,proof=found[0];row['native_shipments']=original['shipments'];row['native_shipping_proof']=proof
